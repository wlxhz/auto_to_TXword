from __future__ import annotations

import logging
from datetime import datetime
from typing import Any

from openpyxl import load_workbook

from tencent_doc import TencentDocClient, _values_equal


def test_mock_backend_exports_excel(tmp_path) -> None:
    client = TencentDocClient(
        {
            "backend": "mock",
            "column_order": ["序号", "姓名", "性别", "首次佩戴日期", "血糖均值"],
            "mock": {"output_file": "output/result.xlsx", "sheet_name": "血糖客户汇总"},
        },
        tmp_path,
        logging.getLogger("test"),
    )
    result = client.update_sheet(
        [{"序号": 1, "姓名": "测试客户", "首次佩戴日期": "2026-09-20", "血糖均值": 6.25}]
    )
    output = tmp_path / "output" / "result.xlsx"
    assert output.exists()
    assert result["records"] == 1
    worksheet = load_workbook(output, read_only=True).active
    assert list(worksheet.values) == [
        ("序号", "姓名", "性别", "首次佩戴日期", "血糖均值"),
        (1, "测试客户", None, datetime(2026, 9, 20, 0, 0), 6.25),
    ]


def test_mock_followup_dates_have_today_red_font_rule(tmp_path) -> None:
    client = TencentDocClient(
        {
            "backend": "mock",
            "column_order": ["首次佩戴日期", "1小时回访", "3天回访", "首次周回访", "取机器回访"],
            "mock": {"output_file": "result.xlsx"},
        },
        tmp_path,
        logging.getLogger("test"),
    )
    client.update_sheet([{
        "首次佩戴日期": "2026-10-01",
        "1小时回访": "2026-10-01",
        "3天回访": "2026-10-04",
        "首次周回访": "2026-10-10",
        "取机器回访": "2026-10-14",
    }])
    worksheet = load_workbook(tmp_path / "result.xlsx").active
    rules = list(worksheet.conditional_formatting)
    assert len(rules) == 1
    assert str(rules[0].sqref) == "B2:E2"
    assert worksheet.conditional_formatting[rules[0]][0].formula == [
        "AND(ISNUMBER(B2),INT(B2)=TODAY())"
    ]


class FakeMCPClient:
    def __init__(self) -> None:
        self.calls: list[tuple[str, dict[str, Any]]] = []

    def call_tool(self, name: str, arguments: dict[str, Any]) -> dict[str, Any]:
        self.calls.append((name, arguments))
        if name == "smartsheet.list_records":
            return {
                "records": [
                    {
                        "record_id": "r1",
                        "field_values": {
                            "姓名": [{"text": "测试客户甲", "type": "text"}],
                            "首次佩戴日期": "1789833600000",
                            "序号": 1,
                            "血糖均值": 6.0,
                            "血糖最高值": 9.0,
                            "血糖最低值": 4.0,
                            "1小时回访": [{"text": "人工已回访", "type": "text"}],
                        },
                    },
                    {
                        "record_id": "r2",
                        "field_values": {
                            "姓名": [{"text": "测试客户乙", "type": "text"}],
                            "首次佩戴日期": "1789920000000",
                            "序号": 2,
                            "血糖均值": 7.1,
                            "血糖最高值": 12.6,
                            "血糖最低值": 4.2,
                        },
                    },
                    {
                        "record_id": "r3",
                        "field_values": {
                            "姓名": [{"text": "文档独有客户", "type": "text"}],
                            "首次佩戴日期": "1790006400000",
                        },
                    },
                ],
                "has_more": False,
            }
        return {"error": ""}


def test_review_existing_updates_only_changed_managed_fields(tmp_path) -> None:
    client = TencentDocClient({}, tmp_path, logging.getLogger("test"))
    fake = FakeMCPClient()
    source = [
        {
            "field_values": {
                "姓名": [{"text": "测试客户甲", "type": "text"}],
                "首次佩戴日期": "1789833600000",
                "序号": 1,
                "血糖均值": 6.3,
                "血糖最高值": None,
                "血糖最低值": 4.0,
            }
        },
        {
            "field_values": {
                "姓名": [{"text": "测试客户乙", "type": "text"}],
                "首次佩戴日期": "1789920000000",
                "序号": 2,
                "血糖均值": 7.1,
                "血糖最高值": 12.6,
                "血糖最低值": 4.2,
            }
        },
        {
            "field_values": {
                "姓名": [{"text": "服务器新增客户", "type": "text"}],
                "首次佩戴日期": "1790092800000",
                "序号": 4,
                "血糖均值": 5.5,
                "血糖最高值": 8.5,
                "血糖最低值": 3.5,
            }
        },
    ]
    config = {
        "primary_keys": ["姓名", "首次佩戴日期"],
        "primary_key_types": {"姓名": "text", "首次佩戴日期": "date"},
        "date_timezone_offset_hours": 8,
        "managed_fields": ["序号", "姓名", "首次佩戴日期", "血糖均值", "血糖最高值", "血糖最低值"],
        "new_record_policy": "skip",
        "ignore_blank_updates": True,
        "batch_size": 100,
        "page_size": 100,
    }

    result = client._review_existing_smartsheet(fake, "file", "sheet", source, config)

    assert result["updated"] == 1
    assert result["unchanged"] == 1
    assert result["unmatched_source"] == 1
    assert result["unmatched_target"] == 1
    update_calls = [args for name, args in fake.calls if name == "smartsheet.update_records"]
    assert update_calls[0]["records"] == [
        {"record_id": "r1", "field_values": {"血糖均值": 6.3}}
    ]


def test_date_values_compare_by_beijing_calendar_day() -> None:
    assert _values_equal("1790784000000", "2026-10-01", "date", 8)
    assert not _values_equal("2026-10-01", "2026-10-04", "date", 8)


def test_customer_id_migration_adds_new_customer_and_marks_updated_date(tmp_path) -> None:
    class Client:
        def __init__(self):
            self.calls = []

        def call_tool(self, name, arguments):
            self.calls.append((name, arguments))
            if name == "smartsheet.list_records":
                return {"records": [{
                    "record_id": "existing-1",
                    "field_values": {
                        "姓名": [{"text": "客户甲", "type": "text"}],
                        "首次佩戴日期": [{"text": "2026-10-01", "type": "text"}],
                        "血糖均值": [{"text": "6.0", "type": "text"}],
                    },
                }], "has_more": False}
            return {}

    fake = Client()
    client = TencentDocClient({}, tmp_path, logging.getLogger("test"))
    source = [
        {"field_values": {
            "客户ID": [{"text": "uuid-1", "type": "text"}],
            "姓名": [{"text": "客户甲", "type": "text"}],
            "首次佩戴日期": [{"text": "2026-10-01", "type": "text"}],
            "血糖均值": [{"text": "6.5", "type": "text"}],
        }},
        {"field_values": {
            "客户ID": [{"text": "uuid-2", "type": "text"}],
            "姓名": [{"text": "客户乙", "type": "text"}],
            "首次佩戴日期": [{"text": "2026-10-02", "type": "text"}],
            "血糖均值": [{"text": "5.9", "type": "text"}],
        }},
    ]
    config = {
        "primary_keys": ["客户ID"],
        "legacy_primary_keys": ["姓名", "首次佩戴日期"],
        "legacy_primary_key_types": {"首次佩戴日期": "date"},
        "managed_fields": ["客户ID", "姓名", "首次佩戴日期", "血糖均值"],
        "last_updated_field": "最近更新日期",
        "review_date": "2026-10-04",
        "new_record_policy": "add",
        "field_types": {"客户ID": "text", "最近更新日期": "text"},
    }
    result = client._review_existing_smartsheet(fake, "file", "sheet", source, config)
    assert result["updated"] == 1
    assert result["added"] == 1
    updates = [args for name, args in fake.calls if name == "smartsheet.update_records"]
    adds = [args for name, args in fake.calls if name == "smartsheet.add_records"]
    assert updates[0]["records"][0]["field_values"] == {
        "客户ID": [{"text": "uuid-1", "type": "text"}],
        "血糖均值": [{"text": "6.5", "type": "text"}],
        "最近更新日期": [{"text": "2026-10-04", "type": "text"}],
    }
    assert adds[0]["records"][0]["field_values"]["客户ID"][0]["text"] == "uuid-2"
    assert adds[0]["records"][0]["field_values"]["最近更新日期"][0]["text"] == "2026-10-04"
