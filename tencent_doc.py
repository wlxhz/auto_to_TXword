"""腾讯文档适配器。

支持离线 Excel（mock）和腾讯文档官方 Streamable HTTP MCP 两种后端。
"""

from __future__ import annotations

import json
import logging
import math
import re
import unicodedata
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Iterable

import requests
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


class TencentDocError(RuntimeError):
    """腾讯文档写入失败。"""


class MCPHttpClient:
    """满足本项目所需的轻量 Streamable HTTP MCP 客户端。"""

    def __init__(self, endpoint: str, token: str, timeout: float, protocol_version: str):
        if not endpoint:
            raise TencentDocError("tencent_doc.mcp.endpoint 不能为空")
        if not token:
            raise TencentDocError("缺少腾讯文档 Token，请设置 TENCENT_DOCS_TOKEN")
        self.endpoint = endpoint
        self.timeout = timeout
        self.protocol_version = protocol_version
        self.session = requests.Session()
        self.session.headers.update(
            {
                "Authorization": token,
                "Accept": "application/json, text/event-stream",
                "Content-Type": "application/json",
            }
        )
        self._request_id = 0
        self._initialized = False

    def _next_id(self) -> int:
        self._request_id += 1
        return self._request_id

    @staticmethod
    def _decode_response(response: requests.Response) -> dict[str, Any]:
        content_type = response.headers.get("Content-Type", "")
        if "text/event-stream" in content_type:
            messages = []
            for line in response.text.splitlines():
                if line.startswith("data:"):
                    try:
                        messages.append(json.loads(line[5:].strip()))
                    except json.JSONDecodeError:
                        continue
            if not messages:
                raise TencentDocError("腾讯文档 MCP 返回了空的 SSE 响应")
            return messages[-1]
        try:
            return response.json()
        except ValueError as exc:
            raise TencentDocError("腾讯文档 MCP 返回了无法解析的响应") from exc

    def _post(self, payload: dict[str, Any], expect_response: bool = True) -> dict[str, Any]:
        try:
            response = self.session.post(self.endpoint, json=payload, timeout=self.timeout)
            response.raise_for_status()
        except requests.RequestException as exc:
            raise TencentDocError(f"腾讯文档 MCP 请求失败：{exc}") from exc

        session_id = response.headers.get("Mcp-Session-Id")
        if session_id:
            self.session.headers["Mcp-Session-Id"] = session_id
        if not expect_response or response.status_code == 202 or not response.content:
            return {}
        message = self._decode_response(response)
        if "error" in message:
            raise TencentDocError(f"腾讯文档 MCP 协议错误：{message['error']}")
        return message.get("result", message)

    def initialize(self) -> None:
        if self._initialized:
            return
        self._post(
            {
                "jsonrpc": "2.0",
                "id": self._next_id(),
                "method": "initialize",
                "params": {
                    "protocolVersion": self.protocol_version,
                    "capabilities": {},
                    "clientInfo": {"name": "nutrition-ledger-automation", "version": "1.0.0"},
                },
            }
        )
        self._post(
            {"jsonrpc": "2.0", "method": "notifications/initialized", "params": {}},
            expect_response=False,
        )
        self._initialized = True

    def call_tool(self, name: str, arguments: dict[str, Any]) -> dict[str, Any]:
        self.initialize()
        result = self._post(
            {
                "jsonrpc": "2.0",
                "id": self._next_id(),
                "method": "tools/call",
                "params": {"name": name, "arguments": arguments},
            }
        )
        if result.get("isError"):
            raise TencentDocError(f"腾讯文档工具 {name} 执行失败：{_content_text(result)}")
        structured = result.get("structuredContent")
        if isinstance(structured, dict):
            payload = structured
        else:
            text = _content_text(result)
            try:
                payload = json.loads(text) if text else result
            except json.JSONDecodeError:
                payload = {"message": text}
        if isinstance(payload, dict) and payload.get("error"):
            raise TencentDocError(f"腾讯文档工具 {name} 返回错误：{payload['error']}")
        return payload


def _content_text(result: dict[str, Any]) -> str:
    return "\n".join(
        item.get("text", "")
        for item in result.get("content", [])
        if isinstance(item, dict) and item.get("type") == "text"
    )


def _chunks(items: list[Any], size: int) -> Iterable[list[Any]]:
    for index in range(0, len(items), size):
        yield items[index : index + size]


def _format_value(value: Any, field_type: str, timezone_offset_hours: float = 8) -> Any:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return None
    if field_type == "text":
        return [{"text": str(value), "type": "text"}]
    if field_type == "number":
        return float(value)
    if field_type == "integer":
        return int(value)
    if field_type == "checkbox":
        return bool(value)
    if field_type == "date":
        if isinstance(value, (datetime, date)):
            dt = datetime.combine(value, datetime.min.time()) if isinstance(value, date) and not isinstance(value, datetime) else value
        else:
            try:
                dt = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
            except ValueError as exc:
                raise TencentDocError(f"无法解析日期值：{value!r}") from exc
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone(timedelta(hours=timezone_offset_hours)))
        return str(int(dt.timestamp() * 1000))
    return value


class TencentDocClient:
    def __init__(self, config: dict[str, Any], base_dir: Path, logger: logging.Logger):
        self.config = config
        self.base_dir = base_dir
        self.logger = logger

    def update_sheet(self, records: list[dict[str, Any]]) -> dict[str, Any]:
        backend = self.config.get("backend", "mock")
        if backend == "mock":
            return self._export_excel(records)
        if backend != "mcp":
            raise TencentDocError(f"未知腾讯文档后端：{backend}")
        return self._write_mcp(records)

    def _export_excel(self, records: list[dict[str, Any]]) -> dict[str, Any]:
        output = Path(self.config.get("mock", {}).get("output_file", "output/tencent_doc_preview.xlsx"))
        if not output.is_absolute():
            output = self.base_dir / output
        output.parent.mkdir(parents=True, exist_ok=True)
        sheet_name = self.config.get("mock", {}).get("sheet_name", "数据")
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = sheet_name
        columns = list(self.config.get("column_order", []))
        if not columns:
            columns = list(records[0].keys()) if records else []
        if columns:
            worksheet.append(columns)
            for record in records:
                worksheet.append([_excel_cell_value(column, record.get(column)) for column in columns])
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            thin = Side(style="thin", color="D9D9D9")
            for cell in worksheet[1]:
                cell.font = Font(name="Arial", size=11, bold=True)
                cell.fill = PatternFill("solid", fgColor="FFFFFF")
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            for row in worksheet.iter_rows(min_row=2):
                for cell in row:
                    cell.font = Font(name="Arial", size=11)
                    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
                    cell.alignment = Alignment(vertical="center")
            for column_name in ("血糖均值", "血糖最高值", "血糖最低值"):
                if column_name in columns:
                    index = columns.index(column_name) + 1
                    for cell in worksheet.iter_cols(min_col=index, max_col=index, min_row=2):
                        for item in cell:
                            item.number_format = "0.00"
            for column_name in ("首次佩戴日期", "1小时回访", "3天回访", "首次周回访", "取机器回访"):
                if column_name in columns:
                    index = columns.index(column_name) + 1
                    for cell in worksheet.iter_cols(min_col=index, max_col=index, min_row=2):
                        for item in cell:
                            item.number_format = "yyyy-mm-dd"
            # 按系统日期动态标红计划回访日期；仅作用于当前输出的数据行。
            if records and all(name in columns for name in ("1小时回访", "3天回访", "首次周回访", "取机器回访")):
                first = columns.index("1小时回访") + 1
                last = columns.index("取机器回访") + 1
                top_left = worksheet.cell(2, first).coordinate
                target = f"{top_left}:{worksheet.cell(len(records) + 1, last).coordinate}"
                worksheet.conditional_formatting.add(
                    target,
                    FormulaRule(
                        formula=[f"AND(ISNUMBER({top_left}),INT({top_left})=TODAY())"],
                        font=Font(color="FF0000"),
                    ),
                )
            if "电话" in columns:
                index = columns.index("电话") + 1
                for cell in worksheet.iter_cols(min_col=index, max_col=index, min_row=2):
                    for item in cell:
                        item.number_format = "@"
            widths = {
                "序号": 8,
                "姓名": 22,
                "性别": 10,
                "电话": 16,
                "是否是糖尿病": 17,
                "是否服药": 13,
                "首次佩戴日期": 16,
                "1小时回访": 14,
                "3天回访": 14,
                "首次周回访": 15,
                "取机器回访": 15,
                "血糖均值": 13,
                "血糖最高值": 13,
                "血糖最低值": 13,
                "14天血糖达标情况": 22,
                "客户ID": 38,
                "最近更新日期": 18,
            }
            for index, column_name in enumerate(columns, start=1):
                worksheet.column_dimensions[worksheet.cell(1, index).column_letter].width = widths.get(column_name, 14)
        workbook.save(output)
        self.logger.info("已生成腾讯文档写入预览：%s", output)
        return {"backend": "mock", "file": str(output), "records": len(records)}

    def _write_mcp(self, records: list[dict[str, Any]]) -> dict[str, Any]:
        mcp = self.config.get("mcp", {})
        client = MCPHttpClient(
            endpoint=mcp.get("endpoint", ""),
            token=mcp.get("token", ""),
            timeout=float(mcp.get("timeout_seconds", 30)),
            protocol_version=mcp.get("protocol_version", "2025-03-26"),
        )
        document_type = self.config.get("document_type", "smartsheet")
        if document_type == "sheet":
            from sheet_sync import sync_sheet
            return sync_sheet(client, records, self.config["sheet"], self.base_dir)
        if document_type == "smartsheet":
            return self._write_smartsheet(client, records)
        if document_type == "excel":
            return self._append_excel(client, records)
        raise TencentDocError(f"不支持的 document_type：{document_type}")

    def _write_smartsheet(self, client: MCPHttpClient, records: list[dict[str, Any]]) -> dict[str, Any]:
        sheet = self.config.get("smartsheet", {})
        if sheet.get("typed_api"):
            from smart_sync import sync_smart
            return sync_smart(client, records, sheet, self.base_dir)
        file_id, sheet_id = sheet.get("file_id", ""), sheet.get("sheet_id", "")
        if not file_id or not sheet_id:
            raise TencentDocError("smartsheet.file_id 和 smartsheet.sheet_id 不能为空")
        field_types = sheet.get("field_types", {})
        timezone_offset_hours = float(sheet.get("date_timezone_offset_hours", 8))
        formatted = [
            {
                "field_values": {
                    key: _format_value(value, field_types.get(key, "raw"), timezone_offset_hours)
                    for key, value in record.items()
                }
            }
            for record in records
        ]
        strategy = sheet.get("sync_strategy", "append")
        if strategy == "review_existing":
            return self._review_existing_smartsheet(client, file_id, sheet_id, formatted, sheet)
        if strategy == "upsert":
            return self._upsert_smartsheet(client, file_id, sheet_id, formatted, sheet)
        if strategy != "append":
            raise TencentDocError(f"不支持的智能表同步策略：{strategy}")

        count = 0
        for batch in _chunks(formatted, int(sheet.get("batch_size", 100))):
            client.call_tool(
                "smartsheet.add_records",
                {"file_id": file_id, "sheet_id": sheet_id, "records": batch},
            )
            count += len(batch)
        return {"backend": "mcp", "mode": "smartsheet_append", "records": count}

    def _review_existing_smartsheet(
        self,
        client: MCPHttpClient,
        file_id: str,
        sheet_id: str,
        records: list[dict[str, Any]],
        config: dict[str, Any],
    ) -> dict[str, Any]:
        """仅比较并更新已存在记录的受管字段，不触碰人工维护列。"""
        primary_keys = list(config.get("primary_keys", []))
        if not primary_keys and config.get("primary_key"):
            primary_keys = [config["primary_key"]]
        if not primary_keys:
            raise TencentDocError("review_existing 模式必须配置 smartsheet.primary_keys")

        managed_fields = list(config.get("managed_fields", []))
        if not managed_fields:
            managed_fields = list(records[0].get("field_values", {}).keys()) if records else primary_keys
        for key in primary_keys:
            if key not in managed_fields:
                managed_fields.append(key)

        legacy_keys = list(config.get("legacy_primary_keys", []))
        lookup_fields = list(dict.fromkeys([*managed_fields, *legacy_keys]))

        existing_items = self._list_smartsheet_records(
            client, file_id, sheet_id, lookup_fields, int(config.get("page_size", 100))
        )
        existing_by_key: dict[tuple[str, ...], dict[str, Any]] = {}
        legacy_by_key: dict[tuple[str, ...], dict[str, Any]] = {}
        key_types = dict(config.get("primary_key_types", {}))
        legacy_key_types = dict(config.get("legacy_primary_key_types", {}))
        timezone_offset_hours = float(config.get("date_timezone_offset_hours", 8))
        duplicate_target_keys = 0
        for item in existing_items:
            values = item.get("field_values", {})
            key = _record_key(values, primary_keys, key_types, timezone_offset_hours)
            if all(key):
                if key in existing_by_key:
                    duplicate_target_keys += 1
                else:
                    existing_by_key[key] = item
            elif legacy_keys:
                legacy_key = _record_key(values, legacy_keys, legacy_key_types, timezone_offset_hours)
                if not all(legacy_key) or legacy_key in legacy_by_key:
                    duplicate_target_keys += 1
                else:
                    legacy_by_key[legacy_key] = item
            else:
                duplicate_target_keys += 1
        if duplicate_target_keys:
            raise TencentDocError(f"腾讯文档中发现 {duplicate_target_keys} 个重复组合键，请先处理后再同步")

        ignore_blank = bool(config.get("ignore_blank_updates", True))
        field_types = dict(config.get("field_types", {}))
        last_updated_field = config.get("last_updated_field", "")
        activity_fields = set(config.get("activity_fields", managed_fields))
        review_date = config.get("review_date", date.today().isoformat())
        new_policy = config.get("new_record_policy", "skip")
        updates: list[dict[str, Any]] = []
        adds: list[dict[str, Any]] = []
        unchanged = 0
        unmatched_source = 0
        seen_keys: set[tuple[str, ...]] = set()
        matched_record_ids: set[str] = set()
        changed_field_counts: dict[str, int] = {}

        for source_item in records:
            source_values = source_item.get("field_values", {})
            source_key = _record_key(source_values, primary_keys, key_types, timezone_offset_hours)
            if any(part == "" for part in source_key):
                raise TencentDocError(f"服务器数据存在空组合键：{' + '.join(primary_keys)}")
            if source_key in seen_keys:
                raise TencentDocError(f"服务器数据存在重复组合键：{' + '.join(primary_keys)}")
            seen_keys.add(source_key)
            target_item = existing_by_key.get(source_key)
            migrated_identity = False
            if target_item is None and legacy_keys:
                legacy_key = _record_key(source_values, legacy_keys, legacy_key_types, timezone_offset_hours)
                if all(legacy_key):
                    target_item = legacy_by_key.get(legacy_key)
                    migrated_identity = target_item is not None
            if not target_item:
                unmatched_source += 1
                if new_policy == "add":
                    added_values = _select_fields(source_values, managed_fields, ignore_blank)
                    if last_updated_field:
                        added_values[last_updated_field] = _format_value(review_date, field_types.get(last_updated_field, "text"))
                    adds.append({"field_values": added_values})
                elif new_policy == "error":
                    raise TencentDocError("服务器存在腾讯文档中找不到的记录；new_record_policy=error")
                elif new_policy != "skip":
                    raise TencentDocError(f"不支持的 new_record_policy：{new_policy}")
                continue

            record_id = target_item.get("record_id", "")
            if not record_id or record_id in matched_record_ids:
                raise TencentDocError("多条服务器记录匹配同一腾讯文档行，请核查客户ID和旧组合键")
            matched_record_ids.add(record_id)

            target_values = target_item.get("field_values", {})
            changed: dict[str, Any] = {}
            if migrated_identity:
                for field in primary_keys:
                    changed[field] = source_values[field]
                    changed_field_counts[field] = changed_field_counts.get(field, 0) + 1
            for field in managed_fields:
                if field in primary_keys or field not in source_values:
                    continue
                source_value = source_values[field]
                if ignore_blank and _is_blank_value(source_value):
                    continue
                if not _values_equal(source_value, target_values.get(field), field_types.get(field), timezone_offset_hours):
                    changed[field] = source_value
                    changed_field_counts[field] = changed_field_counts.get(field, 0) + 1
            if changed:
                if last_updated_field and activity_fields.intersection(changed):
                    changed[last_updated_field] = _format_value(review_date, field_types.get(last_updated_field, "text"))
                updates.append({"record_id": record_id, "field_values": changed})
            else:
                unchanged += 1

        batch_size = int(config.get("batch_size", 100))
        for batch in _chunks(updates, batch_size):
            client.call_tool(
                "smartsheet.update_records",
                {"file_id": file_id, "sheet_id": sheet_id, "records": batch},
            )
        for batch in _chunks(adds, batch_size):
            client.call_tool(
                "smartsheet.add_records",
                {"file_id": file_id, "sheet_id": sheet_id, "records": batch},
            )

        result = {
            "backend": "mcp",
            "mode": "smartsheet_review_existing",
            "source_records": len(records),
            "target_records": len(existing_items),
            "updated": len(updates),
            "unchanged": unchanged,
            "added": len(adds),
            "unmatched_source": unmatched_source,
            "unmatched_target": len(existing_items) - len(matched_record_ids),
            "changed_fields": changed_field_counts,
        }
        self.logger.info("腾讯文档每日审查完成：%s", result)
        return result

    @staticmethod
    def _list_smartsheet_records(
        client: MCPHttpClient,
        file_id: str,
        sheet_id: str,
        fields: list[str],
        limit: int,
    ) -> list[dict[str, Any]]:
        records: list[dict[str, Any]] = []
        offset = 0
        while True:
            result = client.call_tool(
                "smartsheet.list_records",
                {
                    "file_id": file_id,
                    "sheet_id": sheet_id,
                    "field_titles": fields,
                    "offset": offset,
                    "limit": limit,
                },
            )
            records.extend(result.get("records", []))
            if not result.get("has_more"):
                return records
            offset = int(result.get("next", offset + limit))

    def _upsert_smartsheet(
        self,
        client: MCPHttpClient,
        file_id: str,
        sheet_id: str,
        records: list[dict[str, Any]],
        config: dict[str, Any],
    ) -> dict[str, Any]:
        primary_key = config.get("primary_key", "")
        if not primary_key:
            raise TencentDocError("upsert 模式必须配置 smartsheet.primary_key")
        existing: dict[str, str] = {}
        offset, limit = 0, int(config.get("page_size", 100))
        while True:
            result = client.call_tool(
                "smartsheet.list_records",
                {
                    "file_id": file_id,
                    "sheet_id": sheet_id,
                    "field_titles": [primary_key],
                    "offset": offset,
                    "limit": limit,
                },
            )
            for item in result.get("records", []):
                key = _plain_value(item.get("field_values", {}).get(primary_key))
                if key is not None:
                    existing[str(key)] = item.get("record_id", "")
            if not result.get("has_more"):
                break
            offset = int(result.get("next", offset + limit))

        adds, updates = [], []
        for item in records:
            key = str(_plain_value(item["field_values"].get(primary_key)))
            record_id = existing.get(key)
            if record_id:
                updates.append({"record_id": record_id, **item})
            else:
                adds.append(item)

        batch_size = int(config.get("batch_size", 100))
        for batch in _chunks(adds, batch_size):
            client.call_tool("smartsheet.add_records", {"file_id": file_id, "sheet_id": sheet_id, "records": batch})
        for batch in _chunks(updates, batch_size):
            client.call_tool("smartsheet.update_records", {"file_id": file_id, "sheet_id": sheet_id, "records": batch})
        return {"backend": "mcp", "mode": "smartsheet_upsert", "added": len(adds), "updated": len(updates)}

    def _append_excel(self, client: MCPHttpClient, records: list[dict[str, Any]]) -> dict[str, Any]:
        excel = self.config.get("excel", {})
        file_id = excel.get("file_id", "")
        if not file_id:
            raise TencentDocError("excel.file_id 不能为空")
        columns = list(records[0].keys()) if records else []
        rows = [["" if value is None else str(record.get(column, "")) for column in columns] for record in records]
        if excel.get("include_header", False):
            rows.insert(0, columns)
        arguments: dict[str, Any] = {
            "file_id": file_id,
            "string_matrix": {"texts": {"rows": [{"values": row} for row in rows]}},
        }
        if excel.get("sheet_id"):
            arguments["sheet_id"] = excel["sheet_id"]
        arguments.update(excel.get("extra_arguments", {}))
        client.call_tool(excel.get("tool_name", "batch_update_sheet_range"), arguments)
        return {"backend": "mcp", "mode": "excel_append", "records": len(records)}


def _plain_value(value: Any) -> Any:
    if isinstance(value, list) and value:
        first = value[0]
        if isinstance(first, dict):
            return first.get("text", first.get("value"))
        return first
    if isinstance(value, dict):
        return value.get("text", value.get("value"))
    return value


def _record_key(
    values: dict[str, Any],
    fields: list[str],
    field_types: dict[str, str] | None = None,
    timezone_offset_hours: float = 8,
) -> tuple[str, ...]:
    types = field_types or {}
    return tuple(
        _normalize_date_key(_plain_value(values.get(field)), timezone_offset_hours)
        if types.get(field) == "date"
        else _normalize_key_value(_plain_value(values.get(field)))
        for field in fields
    )


def _normalize_key_value(value: Any) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value))
    text = re.sub(r"[\u200b\u200c\u200d\u2060\ufeff]", "", text)
    return " ".join(text.split()).casefold()


def _normalize_date_key(value: Any, timezone_offset_hours: float) -> str:
    if value is None or str(value).strip() == "":
        return ""
    zone = timezone(timedelta(hours=timezone_offset_hours))
    if isinstance(value, (int, float)) or (isinstance(value, str) and value.strip().isdigit()):
        timestamp = float(value)
        if abs(timestamp) > 10_000_000_000:
            timestamp /= 1000
        return datetime.fromtimestamp(timestamp, zone).date().isoformat()
    try:
        parsed = datetime.fromisoformat(str(value).strip().replace("Z", "+00:00"))
        if parsed.tzinfo is not None:
            parsed = parsed.astimezone(zone)
        return parsed.date().isoformat()
    except ValueError:
        return _normalize_key_value(value)


def _is_blank_value(value: Any) -> bool:
    plain = _plain_value(value)
    return plain is None or (isinstance(plain, str) and plain.strip() == "")


def _values_equal(left: Any, right: Any, field_type: str | None = None, timezone_offset_hours: float = 8) -> bool:
    left_plain, right_plain = _plain_value(left), _plain_value(right)
    if _is_blank_value(left) and _is_blank_value(right):
        return True
    if field_type == "date":
        return _normalize_date_key(left_plain, timezone_offset_hours) == _normalize_date_key(right_plain, timezone_offset_hours)
    try:
        return math.isclose(float(left_plain), float(right_plain), rel_tol=1e-9, abs_tol=1e-9)
    except (TypeError, ValueError):
        return _normalize_key_value(left_plain) == _normalize_key_value(right_plain)


def _select_fields(values: dict[str, Any], fields: list[str], ignore_blank: bool) -> dict[str, Any]:
    return {
        field: values[field]
        for field in fields
        if field in values and not (ignore_blank and _is_blank_value(values[field]))
    }


def _excel_cell_value(column: str, value: Any) -> Any:
    if column in ("首次佩戴日期", "1小时回访", "3天回访", "首次周回访", "取机器回访") and isinstance(value, str):
        try:
            return datetime.fromisoformat(value.replace("Z", "+00:00")).date()
        except ValueError:
            return value
    return value


def update_tencent_doc(
    data: list[dict[str, Any]],
    config: dict[str, Any],
    base_dir: Path,
    logger: logging.Logger,
) -> dict[str, Any]:
    """兼容需求文档中的函数式入口。"""
    return TencentDocClient(config, base_dir, logger).update_sheet(data)
